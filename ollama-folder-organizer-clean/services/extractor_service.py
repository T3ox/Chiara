import base64
import io
import os
import re

from constants_py import TRUNCATION_LIMITS


class FileExtractor:
    @staticmethod
    def extract_text_from_docx(file_path: str, snippet: bool = False) -> str:
        from docx import Document

        doc = Document(file_path)
        text = "\n".join(para.text for para in doc.paragraphs)

        if snippet:
            return text[: TRUNCATION_LIMITS["DOCX_SNIPPET_CHARS"]]
        return text[: TRUNCATION_LIMITS["MAX_TEXT_CHARS"]]

    @staticmethod
    def extract_text_from_xlsx(file_path: str) -> str:
        import openpyxl

        wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
        try:
            output = f"Workbook: {os.path.basename(file_path)}\n"

            for sheet_name in wb.sheetnames[: TRUNCATION_LIMITS["XLSX_MAX_SHEETS"]]:
                output += f"\n[SHEET] {sheet_name}\n"
                sheet = wb[sheet_name]

                rows_processed = 0
                for row in sheet.iter_rows(values_only=True):
                    if rows_processed >= TRUNCATION_LIMITS["XLSX_MAX_ROWS_PER_SHEET"]:
                        break

                    row_str = " | ".join(str(cell) if cell is not None else "" for cell in row)
                    if row_str.strip():
                        output += row_str + "\n"
                        rows_processed += 1

            return output[: TRUNCATION_LIMITS["MAX_TEXT_CHARS"]]
        finally:
            wb.close()

    @staticmethod
    def extract_text_from_pptx(file_path: str) -> str:
        from pptx import Presentation

        presentation = Presentation(file_path)
        output = f"Presentation: {os.path.basename(file_path)}\n"

        for slide_index, slide in enumerate(presentation.slides, start=1):
            if slide_index > TRUNCATION_LIMITS["PPTX_MAX_SLIDES"]:
                break

            slide_chunks = []

            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_chunks.append(FileExtractor._normalize_text(shape.text))

                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        row_text = " | ".join(
                            FileExtractor._normalize_text(cell.text)
                            for cell in row.cells
                            if cell.text.strip()
                        )
                        if row_text:
                            slide_chunks.append(row_text)

            if slide_chunks:
                unique_chunks = list(dict.fromkeys(chunk for chunk in slide_chunks if chunk))
                output += f"\n[SLIDE {slide_index}]\n" + "\n".join(unique_chunks) + "\n"

        return output[: TRUNCATION_LIMITS["PPTX_MAX_TEXT_CHARS"]]

    @staticmethod
    def extract_text_from_plain_file(file_path: str) -> str:
        lines = []
        with open(file_path, "r", encoding="utf-8", errors="ignore") as file:
            for index, line in enumerate(file):
                if index >= TRUNCATION_LIMITS["TEXT_MAX_LINES"]:
                    break
                lines.append(line.rstrip())

        return "\n".join(lines)[: TRUNCATION_LIMITS["MAX_TEXT_CHARS"]]

    @staticmethod
    def _normalize_text(text: str) -> str:
        return re.sub(r"\s+", " ", text.replace("\v", " ")).strip()

    @staticmethod
    def extract_text_from_pdf(file_path: str) -> str:
        import fitz

        with fitz.open(file_path) as doc:
            chunks = []
            max_pages = min(doc.page_count, TRUNCATION_LIMITS["PDF_MAX_PAGES_TEXT"])

            for index in range(max_pages):
                page = doc.load_page(index)
                text = page.get_text("text").strip()
                if text:
                    chunks.append(f"[PAGE {index + 1}]\n{text}")

            return "\n\n".join(chunks)[: TRUNCATION_LIMITS["MAX_TEXT_CHARS"]]

    @staticmethod
    def get_pdf_page_count(file_path: str) -> int:
        import fitz

        with fitz.open(file_path) as doc:
            return doc.page_count

    @staticmethod
    def get_pdf_thumbnail(file_path: str) -> str:
        import fitz

        with fitz.open(file_path) as doc:
            page = doc.load_page(0)
            pix = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5))
            img_data = pix.tobytes("jpeg")
            return base64.b64encode(img_data).decode("utf-8")

    @staticmethod
    def resize_image(file_path: str) -> str:
        from PIL import Image

        with Image.open(file_path) as img:
            if img.mode in ("RGBA", "P", "CMYK"):
                img = img.convert("RGB")

            width, height = img.size
            max_side = TRUNCATION_LIMITS["IMAGE_MAX_SIDE"]

            if width > height and width > max_side:
                height = int(height * (max_side / width))
                width = max_side
            elif height >= width and height > max_side:
                width = int(width * (max_side / height))
                height = max_side

            img = img.resize((width, height), Image.Resampling.LANCZOS)
            buffered = io.BytesIO()
            img.save(buffered, format="JPEG", quality=80)
            return base64.b64encode(buffered.getvalue()).decode("utf-8")
