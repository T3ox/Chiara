import base64
import io
import textwrap
import logging
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
import fitz  # PyMuPDF
import openpyxl
import docx
from pptx import Presentation
from bs4 import BeautifulSoup

logger = logging.getLogger(__name__)

class ScreenshotGenerator:
    OUTPUT_SIZE = 1024
    BACKGROUND = "white"
    TEXT_COLOR = "black"

    @staticmethod
    def get_font(size=22):
        try:
            # Trying to load arial, otherwise fallback to default
            return ImageFont.truetype("arial.ttf", size)
        except:
            return ImageFont.load_default()

    @staticmethod
    def fit_to_square(img: Image.Image, size=1024) -> Image.Image:
        img = img.convert("RGB")
        img.thumbnail((size, size), Image.LANCZOS)

        canvas = Image.new("RGB", (size, size), ScreenshotGenerator.BACKGROUND)
        x = (size - img.width) // 2
        y = (size - img.height) // 2
        canvas.paste(img, (x, y))
        return canvas

    @staticmethod
    def render_text_to_image(text: str, title="Anteprima", size=1024) -> Image.Image:
        img = Image.new("RGB", (size, size), ScreenshotGenerator.BACKGROUND)
        draw = ImageDraw.Draw(img)

        title_font = ScreenshotGenerator.get_font(28)
        body_font = ScreenshotGenerator.get_font(20)

        margin = 40
        y = 35

        draw.text((margin, y), title, fill=ScreenshotGenerator.TEXT_COLOR, font=title_font)
        y += 55

        lines = []
        for paragraph in text.splitlines():
            wrapped = textwrap.wrap(paragraph, width=85)
            lines.extend(wrapped if wrapped else [""])

        max_lines = 42

        for line in lines[:max_lines]:
            draw.text((margin, y), line, fill=ScreenshotGenerator.TEXT_COLOR, font=body_font)
            y += 22

        if len(lines) > max_lines:
            draw.text((margin, size - 45), "... contenuto tagliato", fill=ScreenshotGenerator.TEXT_COLOR, font=body_font)

        return img

    @staticmethod
    def image_preview(path: Path) -> Image.Image:
        img = Image.open(path)
        return ScreenshotGenerator.fit_to_square(img, ScreenshotGenerator.OUTPUT_SIZE)

    @staticmethod
    def pdf_preview(path: Path) -> Image.Image:
        doc = fitz.open(path)
        page = doc[0]
        pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
        img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
        return ScreenshotGenerator.fit_to_square(img, ScreenshotGenerator.OUTPUT_SIZE)

    @staticmethod
    def excel_preview(path: Path) -> Image.Image:
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb.active
        rows = []
        for row in ws.iter_rows(max_row=35, max_col=8, values_only=True):
            values = [str(cell) if cell is not None else "" for cell in row]
            rows.append(" | ".join(values))
        text = "\n".join(rows)
        return ScreenshotGenerator.render_text_to_image(text, title=f"Excel: {ws.title}")

    @staticmethod
    def docx_preview(path: Path) -> Image.Image:
        document = docx.Document(path)
        paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
        text = "\n".join(paragraphs[:80])
        return ScreenshotGenerator.render_text_to_image(text, title="Documento Word")

    @staticmethod
    def pptx_preview(path: Path) -> Image.Image:
        prs = Presentation(path)
        content = []

        for i, slide in enumerate(prs.slides, start=1):
            if i > 10:
                break
            content.append(f"--- SLIDE {i} ---")
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text = shape.text.strip()
                    if text:
                        content.append(text)
            content.append("")

        text = "\n".join(content)
        return ScreenshotGenerator.render_text_to_image(
            text[:7000],
            title=f"Presentazione: {path.name}"
        )

    @staticmethod
    def html_preview(path: Path) -> Image.Image:
        html = path.read_text(encoding="utf-8", errors="ignore")
        soup = BeautifulSoup(html, "html.parser")
        text = soup.get_text(separator="\n")
        return ScreenshotGenerator.render_text_to_image(text, title="HTML")

    @staticmethod
    def text_preview(path: Path) -> Image.Image:
        text = path.read_text(encoding="utf-8", errors="ignore")
        return ScreenshotGenerator.render_text_to_image(text[:6000], title=f"Testo: {path.name}")

    def generate_screenshot_b64(self, file_path: str) -> str:
        path = Path(file_path)
        ext = path.suffix.lower()

        image_exts = [".jpg", ".jpeg", ".png", ".webp", ".bmp", ".tif", ".tiff"]
        text_exts = [".txt", ".csv", ".md", ".json", ".py", ".js", ".html", ".css", ".xml"]

        try:
            if ext in image_exts:
                img = self.image_preview(path)
            elif ext == ".pdf":
                img = self.pdf_preview(path)
            elif ext in [".xlsx", ".xlsm"]:
                img = self.excel_preview(path)
            elif ext == ".docx":
                img = self.docx_preview(path)
            elif ext == ".pptx":
                img = self.pptx_preview(path)
            elif ext in [".html", ".htm"]:
                img = self.html_preview(path)
            elif ext in text_exts:
                img = self.text_preview(path)
            else:
                logger.warning(f"Unsupported format for screenshot: {ext}")
                return None

            buffered = io.BytesIO()
            img.save(buffered, format="PNG")
            return base64.b64encode(buffered.getvalue()).decode('utf-8')
        except Exception as e:
            logger.error(f"Error generating screenshot for {file_path}: {e}")
            return None
